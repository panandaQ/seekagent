from __future__ import annotations

from pathlib import Path
from typing import TypedDict

try:
    from openpyxl import load_workbook
    from openpyxl.workbook import Workbook
    from openpyxl.worksheet.worksheet import Worksheet
except ImportError as exc:  # pragma: no cover
    raise RuntimeError("缺少依赖 openpyxl，请先执行: pip install openpyxl") from exc

from config import settings
from models import ExpertInfo
from utils import clean_text


AFFILIATION_HEADER_CANDIDATES = ("单位", "工作单位", "所属单位", "机构", "院校")
STATUS_DONE_VALUES = {"成功"}


class ExpertRow(TypedDict):
    row: int
    name: str
    affiliation: str
    status: str


class ExcelHandler:
    def __init__(self, filepath: str | Path):
        self.filepath: Path = Path(filepath)
        self.workbook: Workbook = load_workbook(self.filepath)
        sheet = self.workbook.active
        if not isinstance(sheet, Worksheet):
            raise RuntimeError("当前工作簿没有可写的工作表")
        self.sheet: Worksheet = sheet
        self.output_columns: dict[str, int] = self._ensure_output_columns()

    def _header_map(self) -> dict[str, int]:
        headers: dict[str, int] = {}
        for col in range(1, self.sheet.max_column + 1):
            value = self.sheet.cell(row=1, column=col).value
            if value is None:
                continue
            headers[str(value).strip()] = col
        return headers

    def _ensure_output_columns(self) -> dict[str, int]:
        """确保所有输出列存在，支持动态字段配置"""
        header_map = self._header_map()
        mapping: dict[str, int] = {}

        # 使用配置中的动态 OUTPUT_HEADERS
        for field_name, header in settings.OUTPUT_HEADERS.items():
            if header in header_map:
                mapping[field_name] = header_map[header]
                continue
            new_col = self.sheet.max_column + 1
            _ = self.sheet.cell(row=1, column=new_col, value=header)
            mapping[field_name] = new_col
            header_map[header] = new_col
        return mapping

    def _find_affiliation_column(self, name_col_idx: int) -> int | None:
        header_map = self._header_map()
        for candidate in AFFILIATION_HEADER_CANDIDATES:
            col_idx = header_map.get(candidate)
            if col_idx and col_idx != name_col_idx:
                return col_idx
        b_header = self.sheet.cell(row=1, column=2).value
        if (
            b_header is not None
            and any(keyword in str(b_header) for keyword in ("单位", "机构"))
            and name_col_idx != 2
        ):
            return 2
        return None

    def read_expert_names(self, name_col: str = "A", start_row: int = 2) -> list[ExpertRow]:
        name_col_idx = ord(name_col.upper()) - ord("A") + 1
        affiliation_col_idx = self._find_affiliation_column(name_col_idx)
        status_col_idx = self.output_columns["status"]
        # 获取输出列中的 affiliation 列索引
        output_affiliation_col_idx = self.output_columns.get("affiliation")

        experts: list[ExpertRow] = []
        for row in range(start_row, self.sheet.max_row + 1):
            raw_name = self.sheet.cell(row=row, column=name_col_idx).value
            name = clean_text(str(raw_name)) if raw_name is not None else None
            if not name:
                continue

            # 优先使用输出列中的 affiliation（上次提取的），其次使用原始输入列
            affiliation = ""
            if output_affiliation_col_idx:
                raw_output_aff = self.sheet.cell(row=row, column=output_affiliation_col_idx).value
                extracted_aff = clean_text(str(raw_output_aff)) if raw_output_aff else None
                if extracted_aff:
                    affiliation = extracted_aff

            if not affiliation and affiliation_col_idx:
                raw_aff = self.sheet.cell(row=row, column=affiliation_col_idx).value
                affiliation = clean_text(str(raw_aff)) or ""

            raw_status = self.sheet.cell(row=row, column=status_col_idx).value
            status = clean_text(str(raw_status)) if raw_status is not None else ""

            experts.append(
                {
                    "row": row,
                    "name": name,
                    "affiliation": affiliation,
                    "status": status or "",
                }
            )
        return experts

    def is_completed(self, status: str | None) -> bool:
        if not status:
            return False
        return status.strip() in STATUS_DONE_VALUES

    def write_result(self, row: int, result: ExpertInfo) -> None:
        """写入结果，支持动态字段"""
        normalized = result.normalized()

        # 遍历所有输出列并写入对应的值
        for field_name in self.output_columns.keys():
            if field_name == "source_urls":
                value = "\n".join(normalized.source_urls)
            elif hasattr(normalized, field_name):
                value = getattr(normalized, field_name)
            else:
                value = normalized[field_name]

            col_idx = self.output_columns[field_name]
            _ = self.sheet.cell(row=row, column=col_idx, value=value)

    def save(self) -> None:
        self.workbook.save(self.filepath)
